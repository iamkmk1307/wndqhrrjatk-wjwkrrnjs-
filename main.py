import os
import re
import io
import base64
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext
from urllib.parse import urlsplit
from copy import copy
import pandas as pd
from openpyxl import load_workbook


def extract_domain(url):
    """
    기준 차단목록과 일일파일을 비교할 때 사용하는 도메인 추출 함수.

    예:
      https://www.t.me/s/funbe_next -> t.me
      t.me/s/toonkor_com            -> t.me

    기준 차단목록 비교에서만 이 값을 사용한다.
    """
    if pd.isna(url):
        return ""

    url = str(url).strip()
    if not url:
        return ""

    if not url.lower().startswith(("http://", "https://")):
        url = "http://" + url

    try:
        parsed = urlsplit(url)
        netloc = parsed.netloc.lower().strip()

        # 사용자정보가 있는 특수 URL 방어
        if "@" in netloc:
            netloc = netloc.rsplit("@", 1)[-1]

        # www. 제거
        if netloc.startswith("www."):
            netloc = netloc[4:]

        # 포트번호 제거
        # 일반적인 host:port 형식 기준
        if ":" in netloc and not netloc.startswith("["):
            netloc = netloc.split(":", 1)[0]

        return netloc

    except Exception:
        return ""


def normalize_full_url(url):
    """
    일일파일 내부 중복 비교 전용.

    도메인만 비교하지 않고 하위 경로/쿼리까지 포함한 전체 URL을 비교한다.

    아래 차이만 무시:
      - http:// 와 https://
      - 맨 앞 www.
      - URL 마지막의 /

    예:
      https://www.t.me/s/funbe_next/
        -> t.me/s/funbe_next

      http://t.me/s/funbe_next
        -> t.me/s/funbe_next

      t.me/s/toonkor_com
        -> t.me/s/toonkor_com

    따라서
      t.me/s/funbe_next != t.me/s/toonkor_com
    이므로 서로 중복이 아니다.
    """
    if pd.isna(url):
        return ""

    value = str(url).strip()
    if not value:
        return ""

    # scheme 제거
    value = re.sub(r"(?i)^https?://", "", value)

    # www. 제거
    value = re.sub(r"(?i)^www\.", "", value)

    # 마지막 /만 제거
    value = re.sub(r"/+$", "", value).strip()

    if not value:
        return ""

    # host 부분만 소문자로 통일하고 path/query/fragment는 그대로 유지
    m = re.match(r"^([^/?#]+)(.*)$", value)
    if not m:
        return value

    host = m.group(1).lower()
    rest = m.group(2)

    return host + rest


def clean_id(value):
    """등록번호/접수번호 값 정리"""
    if pd.isna(value):
        return ""

    value = str(value).strip()

    if value.endswith(".0"):
        value = value[:-2]

    return value



# ============================================================
# 업로드샘플(수정용)(2).xlsx 원본 양식을 코드 내부에 내장
# - 실행파일 옆에 별도 템플릿 파일이 없어도 됨
# - 시트명 / A:Z 제목 / 글꼴 / 정렬 / 테두리 / 너비 / 높이 등
#   샘플 원본 서식을 그대로 사용
# ============================================================
TEMPLATE_XLSX_BASE64 = """
UEsDBBQABgAIAAAAIQBBN4LPbgEAAAQFAAATAAgCW0NvbnRlbnRfVHlwZXNdLnhtbCCiBAIooAACAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAACsVMluwjAQvVfqP0S+Vomhh6qqCBy6HFsk6AeYeJJYJLblGSj8fSdmUVWxCMElUWzPWybzPBit2iZZQkDjbC76WU8kYAunja1y8T39SJ9FgqSsVo2zkIs1oBgN7+8G07UHTLjaYi5qIv8iJRY1tAoz58HyTulCq4g/QyW9KuaqAvnY6z3JwlkCSyl1GGI4eINSLRpK3le8vFEyM1Ykr5tzHVUulPeNKRSxULm0+h9J6srSFKBdsWgZOkMfQGmsAahtMh8MM4YJELExFPIgZ4AGLyPdusq4MgrD2nh8YOtHGLqd4662dV/8O4LRkIxVoE/Vsne5auSPC/OZc/PsNMilrYktylpl7E73Cf54GGV89W8spPMXgc/oIJ4xkPF5vYQIc4YQad0A3rrtEfQcc60C6Anx9FY3F/AX+5QOjtQ4OI+c2gCXd2EXka469QwEgQzsQ3Jo2PaMHPmr2w7dnaJBH+CW8Q4b/gIAAP//AwBQSwMEFAAGAAgAAAAhALVVMCP0AAAATAIAAAsACAJfcmVscy8ucmVscyCiBAIooAACAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAACskk1PwzAMhu9I/IfI99XdkBBCS3dBSLshVH6ASdwPtY2jJBvdvyccEFQagwNHf71+/Mrb3TyN6sgh9uI0rIsSFDsjtnethpf6cXUHKiZylkZxrOHEEXbV9dX2mUdKeSh2vY8qq7iooUvJ3yNG0/FEsRDPLlcaCROlHIYWPZmBWsZNWd5i+K4B1UJT7a2GsLc3oOqTz5t/15am6Q0/iDlM7NKZFchzYmfZrnzIbCH1+RpVU2g5abBinnI6InlfZGzA80SbvxP9fC1OnMhSIjQS+DLPR8cloPV/WrQ08cudecQ3CcOryPDJgosfqN4BAAD//wMAUEsDBBQABgAIAAAAIQCBPpSX8wAAALoCAAAaAAgBeGwvX3JlbHMvd29ya2Jvb2sueG1sLnJlbHMgogQBKKAAAQAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAACsUk1LxDAQvQv+hzB3m3YVEdl0LyLsVesPCMm0KdsmITN+9N8bKrpdWNZLLwNvhnnvzcd29zUO4gMT9cErqIoSBHoTbO87BW/N880DCGLtrR6CRwUTEuzq66vtCw6acxO5PpLILJ4UOOb4KCUZh6OmIkT0udKGNGrOMHUyanPQHcpNWd7LtOSA+oRT7K2CtLe3IJopZuX/uUPb9gafgnkf0fMZCUk8DXkA0ejUISv4wUX2CPK8/GZNec5rwaP6DOUcq0seqjU9fIZ0IIfIRx9/KZJz5aKZu1Xv4XRC+8opv9vyLMv072bkycfV3wAAAP//AwBQSwMEFAAGAAgAAAAhAECyXenWAgAAUAUAAA8AAAB4bC93b3JrYm9vay54bWysVM1rE0EUvwv+D+sQ8JTsRzZpu2RT2iTFgkip/QAJlMnuJDtkd2edmTUpIvTQi7SghygtqBQUvJbioQf9h5rkf/DNbqOpvVT0MjNvPn7vvd/7vaktD6NQe0G4oCx2kVkykEZij/k07rloe2utuIg0IXHs45DFxEX7RKDl+v17tQHj/Q5jfQ0AYuGiQMrE0XXhBSTCosQSEsNJl/EISzB5TxcJJ9gXASEyCnXLMKp6hGmMcgSH3wWDdbvUI03mpRGJZQ7CSYglhC8CmogZWuTdBS7CvJ8mRY9FCUB0aEjlfgaKtMhz1nsx47gTQtpDszJDhuUt6Ih6nAnWlSWA0vMgb+VrGrpp5inXa10akp2cdg0nyRMcKS8h0kIsZMunkvguqoLJBuTGBk+T1ZSGcGratmUgvf6rFBtc80kXp6HcgiLM4OFi1TZMU92EpFZCSXiMJWmwWAKH1+z/K18ZdiNgUB1tkzxPKScgCkVbvQYj9hzcERtYBlrKQxc1nPa2gPTbKYztJhF9yZL2eDSafD2YHl5eXfyYnH0YX4wmnw8nZ4fjo2/T44P25NPb8ZvT6elImxx9mb5/PT0+nnz83p4rDL5d9b8oDfYURzqQlCeSr/8krF5Tst+hZCB+U69MbbhLY58NXARNtD+3HmTbu9SXgYusimmBlvK9R4T2AgkVMq1qJXM+h511CvjIZi3OFGIZVlmbvjsZn59cXZ5Da6puWldiQBp3KCz4up+VWp89B0nQmPhKYQA2Z11D7g3DOCrtrVEljCaWuIMFUcLzcPh0Bg8ZBdT3ifojUP3hzTAePiisFEyn8KxglWv6nAcg8KZ3gPRApWrKgrbtCpABpJOhfCxkNoNAqItemraxsmAs2UWjVa4U7cUlq7hol61iw25arcpCq9larbz6vz0JOnVm35qKMsBcbnHs9eEz3CTdVeBFpa80AvHmYxa1PntV/wkAAP//AwBQSwMEFAAGAAgAAAAhACdvHs3sAwAAJxQAABQAAAB4bC9zaGFyZWRTdHJpbmdzLnhtbORYW2/iVhB+j5T/cOSnVtrgC+ZyKmAfVqq0Uh9W1e5T1QcLzIIa2xQ7tz6hhE1pIGqiBpVNDELK7iapUtUFNiUq/UOc4//QMY66KT4GrbJPIBmEv5kzM994ZnwOqcfb2jraVMtm0dDTnBgROKTqWSNX1F+muRfPv1xLcsi0FD2nrBu6muZ2VJN7nFldSZmmhWCtbqa5gmWVvuB5M1tQNcWMGCVVB0neKGuKBbfll7xZKqtKziyoqqWt85IgxHlNKeocyhobupXmsMShDb34/Yb6xAcSUS6TMouZlJWh3SbpD8hvr1K8lUnxHugLXnz91TRE20N6USHOFek1p2Wkfkl7LDu03qW932nniOGD2lXitMjugJ5eBZydj2jXDqCTcKnddVuBEDx0/z0I3L3guomMXLQDcTs23WvTzlu63wg4+/UVmxO5qdBqwAe5qY77o4CRqk16VWBKb1ts2dNn07i7+4aeNcgvIfRDEtZrkXfX5BaC6AYcvWuP/3S8h3DeEGcJpYmwVIBqtIrZZ2WUN3TraS7NQQlZOyUoUd14Yuh3Jc3x/ysZes9L9CGGgAE8TNKrua3hpwpI/lSGYg8xRBoV99im7QH5w+ulz+jhkFTf0Pbo84dYdZsDUr9Fd718cIKocwkNSY5sRHevwZt7MKSvK4j8/BpcuU0bkZsaNDGsqNDOMbke0XYVbuxx/73bbJGrAZT3ozsrnsg9bbgnsLA1HjoIahkcgvox3T9EMCfIRQORfmXc+8fT6HcR7dQIRAFkoX8+hID8wqWdamR1ZXXlG3p29F94aCu/lZdxPJI1NMRPjNdafgWgRAIn47IM8LjnAEngAKWMJEGKrwl4TZS/DTTQAiSkuB0pGAZc1qYci0W2d35g5iUufEReBMl/Jo8QOTiGn6T+F/174Df9vdnPVptO8gct96zqXXUHpk+4Fm2O3JoTLic/DWbK6Y+juS4Oh/NUSHtEnJO5Wntv5/oazmVMTu3ZjHdbMAdmq1wPYGbPVrmokfNLMvkOPMpFGA07imUYuiTjiGIWFWYXRKNLNh10dctLSlSIz0jKso1Ms2BsWZuiFLGMUrBMEjg2Y1gKC/kSMQ1NhYzIISnBEpalJeuc73RrMxq+z4iJS5aPXG4yXSX2ziuZTMSToRlZzKbRFU1VRRgjIa8baJplK5K8kRUleUZCPmYXugh7EkGMaEp5w/uIYkjrROOx8OEabJ1EUhCjycDx3IcDGzsf9o/Z97buPuwfdgOwf3QNwPHpPblvJMGGk2wYM2EssGE2S8xmidksMZslZrPEbJaYzRKzWWImS1lgspQFJktZuMeShz/6Mv8CAAD//wMAUEsDBBQABgAIAAAAIQA7bTJLwQAAAEIBAAAjAAAAeGwvd29ya3NoZWV0cy9fcmVscy9zaGVldDEueG1sLnJlbHOEj8GKwjAURfcD/kN4e5PWhQxDUzciuFXnA2L62gbbl5D3FP17sxxlwOXlcM/lNpv7PKkbZg6RLNS6AoXkYxdosPB72i2/QbE46twUCS08kGHTLr6aA05OSonHkFgVC7GFUST9GMN+xNmxjgmpkD7m2UmJeTDJ+Ysb0Kyqam3yXwe0L0617yzkfVeDOj1SWf7sjn0fPG6jv85I8s+ESTmQYD6iSDnIRe3ygGJB63f2nmt9DgSmbczL8/YJAAD//wMAUEsDBBQABgAIAAAAIQBAD/I6nQYAAI0aAAATAAAAeGwvdGhlbWUvdGhlbWUxLnhtbOxZW4sbNxR+L/Q/DPPu+DbjyxJvsMd2ts1uErJOSh5lW/YoqxmZkbwbEwIlodBCKRTS0pdC3/JQSgMNNPSlP2YhoU37H3qkGXuktZzNZVPSkjUsM5pPR5/OOfp0O3/hdkSdQ5xwwuKWWz5Xch0cj9iYxNOWe33QLzRchwsUjxFlMW65C8zdC9sffnAebYkQR9iB+jHfQi03FGK2VSzyERQjfo7NcAzfJiyJkIDXZFocJ+gI7Ea0WCmVasUIkdh1YhSB2SuTCRlh5+/Pvnz+8HN3e2m9R6GJWHBZMKLJvrSNjSoKOz4oSwRf8IAmziGiLRcaGrOjAb4tXIciLuBDyy2pP7e4fb6ItrJKVGyoq9Xrq7+sXlZhfFBRbSbT4apRz/O9WntlXwGoWMf16r1ar7aypwBoNIKeplx0m36n2en6GVYDpY8W2916t1o28Jr96hrnti9/Bl6BUvveGr7fD8CLBl6BUrxv8Um9EngGXoFSfG0NXy+1u17dwCtQSEl8sIYu+bVqsOztCjJhdMcKb/pev17JjOcoyIZVdskmJiwWm3ItQrdY0geABFIkSOyIxQxP0AjSOECUDBPi7JJpCIk3QzHjUFyqlPqlKvyXP089KY+gLYy02pIXMOFrRZKPw0cJmYmW+zFYdTXI0ydPju89Pr736/H9+8f3fs7aVqaMejsonur1nj/8+q/vP3X+/OWH5w++SZs+iec6/tlPXzz77fcXmYce5654+u2jZ48fPf3uqz9+fGCx3k7QUIcPSIS5cxkfOddYBB208MfD5NVqDEJEjBooBNsW0z0RGsDLC0RtuA42XXgjAZWxAS/Obxlc98NkLoil5UthZAD3GKMdllgdcEm2pXl4MI+n9saTuY67htChre0AxUaAe/MZyCuxmQxCbNC8SlEs0BTHWDjyGzvA2NK7m4QYft0jo4RxNhHOTeJ0ELG6ZECGRiLllXZIBHFZ2AhCqA3f7N1wOozaet3FhyYShgWiFvIDTA03XkRzgSKbyQGKqO7wXSRCG8n9RTLScT0uINJTTJnTG2PObXWuJNBfLeiXQGHsYd+ji8hEJoIc2GzuIsZ0ZJcdBCGKZlbOJA517Ef8AFIUOVeZsMH3mDlC5DvEAcUbw32DYCPcpwvBdRBXnVKeIPLLPLHE8iJm5nhc0AnCSmVA+w1Jj0h8qr6fUHb/31F2u0afgabbDb+JmrcTYh1TOyc0fBPuP6jcXTSPr2IYLOsz13vhfi/c7v9euDeN5bOX61yhQbzztbpauUcbF+4TQum+WFC8y9XancO8NO5DodpUqJ3laiM3C+Ex2yYYuGmCVB0nYeITIsL9EM1ggV9W29Apz0xPuTNjHNb9qljtiPEJ22r3MI/22Djdr5bLcm+aigdHIi8v+aty2GuIFF2r53uwlXm1q52qvfKSgKz7KiS0xkwSVQuJ+rIQovAiEqpnZ8KiaWHRkOaXoVpGceUKoLaKCiycHFhutVzfS88BYEuFKB7LOKVHAsvoyuCcaaQ3OZPqGQCriGUG5JFuSq4buyd7l6baS0TaIKGlm0lCS8MQjXGWnfrByVnGupmH1KAnXbEcDTmNeuNtxFqKyAltoLGuFDR2jlpurerD4dgIzVruBPb98BjNIHe4XPAiOoXTs5FI0gH/OsoyS7joIh6mDleik6pBRAROHEqiliu7v8oGGisNUdzKFRCEd5ZcE2TlXSMHQTeDjCcTPBJ62LUS6en0FRQ+1QrrV1X99cGyJptDuPfD8ZEzpPPkGoIU8+tl6cAx4XD8U069OSZwnrkSsjz/TkxMmezqB4oqh9JyRGchymYUXcxTuBLRFR31tvKB9pb1GRy67sLhVE6wbzzrnj5VS89popnPmYaqyFnTLqZvb5LXWOWTqMEqlW61beC51jWXWgeJap0lTpl1X2JC0KjljRnUJON1GZaanZWa1M5wQaB5orbBb6s5wuqJ1535od7JrJUTxHJdqRJf3XzodxNseAvEowunwHMquAol3DwkCBZ96TlyKhswRG6LbI0IT848IS33Tslve0HFDwqlht8reFWvVGj47Wqh7fvVcs8vl7qdyl2YWEQYlf301qUPB1F0kd29qPK1+5doedZ2bsSiIlP3K0VFXN2/lCu2+5eBvF9xHQKic6dW6TerzU6t0Ky2+wWv22kUmkGtU+jWgnq33w38RrN/13UOFdhrVwOv1msUauUgKHi1kqTfaBbqXqXS9urtRs9r382WMdDzVD4yX4B7Fa/tfwAAAP//AwBQSwMEFAAGAAgAAAAhAHpVLa+wBAAAfxcAAA0AAAB4bC9zdHlsZXMueG1s1Fjdbts2FL4fsHcQeO9IsiXHMiwXdRwBBbpiQDxgt7RE2UQo0qDo1O4woK9QYLtbgV4M2APsrdbuHXZISZY8J7HrOEiaAIlIHR595//wDF6sMmbdEJlTwUPknjnIIjwWCeWzEP00iVo9ZOUK8wQzwUmI1iRHL4bffzfI1ZqRqzkhygIWPA/RXKlF37bzeE4ynJ+JBeHwJhUywwqWcmbnC0lwkutDGbPbjtO1M0w5Kjj0s/gQJhmW18tFKxbZAis6pYyqteGFrCzuv5pxIfGUAdSV6+HYWrld2a6+YLZ2PpLRWIpcpOoMmNoiTWlMdrEGdmDjuOYEbI/j5Pq20y4EHw5SwVVuxWLJVYhA1wZh/5qLtzzSr8AmqKAaDvJ31g1msOMieziIBRPSUqBskNXscJyRguLzXx++fHxv/fP3p8+//a6JU5xRti5ets3pOZY52K5g2A70nrFcySGjoEe9aWuEFYKppipRmCNHg+gcD6IC0NMsHl/mZyPvA81+Ao03TC5n0xBFkWN+vsIQJ0BxCvff7+pbZj/Qy44SzgRYDhFGGdtkAl8HPWwMB5DkFJE8goVVPk/WCwh5Dvm4CFBDt4d6JvHabfuHH8gFo4lGMbtoJhofWYrqXOWcnQdB0HO7vV4v8Dqu55m8Mi3JKU/IiiQh6nrmmw0xdE45BPIdCM6fHIF3QgRGFWD9qZAJVOGqEnRA88XWcMBIqiDTSTqb6/9KLODvVCgFxWo4SCieCY6ZTtXVieZJqN5QqEOk5lBo/1c1jMVszb9kv5/YoDAg9tMC0grofuJCoCPlKbPRONC/xt/2CbV7Yo9kuwf2ibd74j4ZS+OBK8SEsStttJ/TLX9YpRZfZlGmXkFYQZ+mC3P1CPFUPha2LxbDAWZ0xjPCodATqWis24cYlqSo7at0OLiLLXi5DtNTs+3eyVZ7cFP2QhMNJbg6K369FqxVegJ1uDVuaCdrLcN+xd/CiwVb67ZNN2TlClRYr0Ymyuv1y8o8RZtXW2suJH0HjBr22rWg9VbixYSszOe09u4z53PEn88l5dcTEdGDROjUJoDHb88EDfwQvgfjf0KngWpf5QBIB7cg3rj583T6hsa3ghbUf3jQPixI98Xls3CKY0Hu0eMjaM4L6sJ3l+p2UL1ZZlMiIzOJaKTmLZ99Iqw75eMQrPoy8FhF5VZXeDStb0vyNJH2jJIB3G2qfPtNVjjdld7WM+oBXNmEFW3RwTWtbluh3Ua3djymaYQ2sdE5b/XNm57S0tf4EP374Y8vf75vIJouKYNLbdEk6sla1YBvHbA2QhQDtNuJwGyFpGYiV/ezgC5Z1f28oy8pSs8pTae/wQsKTEiKl0xNNi9DVD//QBK6zABJSfUjvRHKsAhR/fxaXxXdrv4GdIevc7jewX9rKWmIfrkcnQfjy6jd6jmjXsvrEL8V+KNxy/cuRuNxFDht5+LXxsD0AeNSM9yFlt31+jmDoaoshS3BX9V7IWosCvjmGgewm9iDdtd56btOK+o4bsvr4l6r1+34rch32+OuN7r0I7+B3T9yQOvYrlsNaFeu31c0I4zyylaVhZq7YCRY3iOEXVnCrifnw/8AAAD//wMAUEsDBBQABgAIAAAAIQDfUKTkHwcAAFciAAAYAAAAeGwvd29ya3NoZWV0cy9zaGVldDEueG1srFrbcqJKFH0/VecfLN5HA/ESUzFTk3FMzH3QXN8IYqRGhQIymczXn71pkO5FGw05L4rL1Zvdi9W9m4aDr38W89pvL4r9YNkzzPqOUfOWbjDxl88942Y8+LJn1OLEWU6cebD0esabFxtfD//95+A1iH7FM89LahRhGfeMWZKE+41G7M68hRPXg9Bb0j/TIFo4Cf2MnhtxGHnOJG20mDesnZ12Y+H4S0NE2I+2iRFMp77r9QP3ZeEtExEk8uZOQvnHMz+M82gLd5twCyf69RJ+cYNFSCGe/LmfvKVBjdrC3R8+L4PIeZpTv/+YTcfNY6c/SuEXvhsFcTBN6hSuIRIt97nb6DYo0uHBxKcesOy1yJv2jG/m/qPVNBqHB6lAt773GkvHtcR5Gnlzz028CV0no8b6PwXBLyYOCdrhpo1S20Gq/3VUm3hT52We2MHriec/zxIKYlF/uFv7k7e+F7ukJ4Wp73IgN5jT2emztvDZFySH80ec158ks57RZVu8sTIURjQQVAqaUun7VVDNVr1ttQo+WezJi5OBzzkYNfclToLFXcZVYu1mseg7i7VX3+1UCtXMQtF3npZZrxSJGqUdpO88EvUoF2P7zrWzOJ1PZ0RDNM2IvrOMutW6RhdVXGbqxSdVMleWYa9mF7eSTmxTkZVkqYrXzswdxQefyyo3lCk5yqpX87mZe4oPVk6v5E4ztxUffPYakjWF8oVHu/W9agPQzF3KB3liVt2sqFhuVUuy6l7VaFZuVkuy2F5p1npnprJyY/FB3juI8P5cR3N/NnFKhoIebYiQu4glzXPoqNdrQ4TcPJZknvaHIqws097dKzqC1aIh6ktarfpO4hweRMFrjRYAlF8cOrycMPc5CTqgTyG8KGUpR1u2qF5xiG8cI21G9SWmQvv70LIOGr+pOroZ5UhQ6HNF2VEZ38sMU2X0yww4y48yY1eNMSgzmirjuMxoqYyTMqOtMoZlRkdlnJYZeyrjrMzoqoxzjWIg6oWGAqpeaigg65WGArpeaygg7E8NBZS1NRSQdqShgLZjDQXEvREUmntWdjRB3dsyxQJ17zQUUPdeQwHpHjQUkO5RUGhcFgOskK5BI3k1nKlLOJxp9i8NZ16NWmaTFy3vjmxZoRb07YjP1TPE6pXnge8I9AWQrrHTmeIHMgYIHGOTEwSGCJwicIbAOQIXCFwicIXAtQBa6XKZe/sTAVsAvDxZXaZdsMwo4xSijTWtLHD0Dap0i8AdAvcIPCDwKAGKhaiebmshs0vmetdAFGwlRguLAp9JMRACfQFIBkLGAIFjbHKCwBCBUwTOEDhH4AKBSwSuELgWgGQgBGwBqAaCQTfKOJKBNK2aoPQNqnSLwB0C9wLorvz+gMCjBCgGogHwf85B8nhqweR5xOdSLIRAXwCShZAxQOAYm5wgMETgFIEzBM4RuEDgEoErBK4FIFkIAVsAqoXADKOMI1lI06oJSt+gSrcC6KwMcieAwjH3CDwg8CgBioX4TgxWpevK2OY5iIIVcxCU1iM+k2IgBPoCkAyEjAECx9jkBIEhAqcInCFwjsAFApcIXCFwLQDJQAjYAlANBFYYZRzJQJpWTVD6BlW6FYBkIAFIBkLgAYFHCVAMxHcysoHS7beN1UpeYbVwcSpujlRloI9jDae5ZqVG922lDLdZklG7ws2wOrY5KN2v0bBauwQZazjSgkNRkbeiUcVtcqR2RY6wmLE56KYcNZx1OfKdbpUcqV2RI9wf2Bx0U44azrocTd6Hq5IkNyyyhFsUOw27KU0daW2efKPx8WHDW4+rLNuw5LX535KYMLjGOlKzcLfiSt6frJIlNSuyhHWVnQZFLWF4jXWktVny8qiCltSsyBJKt827qCUtS/e/uhXhOi25AlfIUp5h2lAfbJ5/SlnCEBvrSM2iK+oV52m+0vihhoWaMFvbvPm7cZjrSGvHj67sbDNh8t5xkWdpW0RXemCYjTkEdmZtnkrxoTlmu/JoyrWnjbWH/y0lAFPWWEta502l/GyfpVx92lh9eDO8lGXJmzrSuiyVArR9lnL9aWP9MXUFqKSlhtQsIqn7QkoF2jpL3uwvfIn1h/9FLZvoSy2piKRmqdSfLMut9qnkCtTBCsRPG0p5QgEY60itIpKap1KBPpKnXIM6WIP4YcjG+UhHWjfO+dlIMW9+JE+5CnWwCqVhN607dKS1eSpVKHdne/P2ZHpXJi6NeBIvnm2EM3p1IvFdego/DZYJP61n4d9Cenq+DL4Hy+z9C753DZ1n78KJnv1lXJt7U3L7Tp1mskg8sk+PkyBMUSoXT0FCz87zXzN6u8KjvUB+hk9nCpL8RxZ35CUvYS10Qi8a+X/p5DRig8in5/7p6xM9IwyiJHL8hGZgwv9Srs68H/p0Q7pDKlCW1AkZifZ96ks0nJjpywerV0IO/wMAAP//AwBQSwMEFAAGAAgAAAAhALo+iVZEAQAAWQIAABEACAFkb2NQcm9wcy9jb3JlLnhtbCCiBAEooAABAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAHySX0vDMBTF3wW/Q8l7m7YrVUPbgcqeHAhOFN9CcrcFmzQkmd2+vemf1crEx9xz7u+ee0mxPMo6+AJjRaNKlEQxCkCxhgu1K9HrZhXeosA6qjitGwUlOoFFy+r6qmCasMbAs2k0GCfABp6kLGG6RHvnNMHYsj1IaiPvUF7cNkZS559mhzVln3QHOI3jHEtwlFNHcQcM9UREI5KzCakPpu4BnGGoQYJyFidRgn+8Doy0fzb0yswphTtpv9MYd87mbBAn99GKydi2bdQu+hg+f4Lf108v/aqhUN2tGKCq4IwwA9Q1pjpYMAWeFbrj1dS6tb/zVgC/P42ey7rn9LEHGPDAByFD7LPytnh43KxQlcZpGsZ5mOabdEGyjMQ3H93YX/1dsKEgx+H/E/MwvguTrCOmGUnmxDOgKvDFZ6i+AQAA//8DAFBLAwQUAAYACAAAACEADOcSLzMEAAA4EgAAJwAAAHhsL3ByaW50ZXJTZXR0aW5ncy9wcmludGVyU2V0dGluZ3MxLmJpbuxX3W9URRQ/5zezc2e/P1ta6MftQltAKQtWrcjHwgqloohFFGs1Nt4mmpiS+PHsxsQH34zxiSf/ACTGxAeSkqjxYXkkxDcT/wkfTIxZz9x72y2FkLUVpMEzmdyZufec85vzNXMr3/9wY+rm9SsDra+Wf27p1u/f0S2f9C/fXqvdpC6I9Rx+pfmi+o2JKUmX05M2kJFHFwF5knSm4zTZjbB/+I2TjlBD9FzP3lhYurQki7YnfiMMjidhiP7SX1t6u6Q/LzdogZboknSfpmlRnov0Ab1H78j8HL1PH9OHMrpAp2hWnjPS1pPD0PaIZtQhGXLACqyURgImYJZxXhdEp7d2YlWSU5xGBlnkkFeUcm+hELLBg0XSCXIr7F5C5LDmREpkevCUJx9aXzSn72CkSsBZWCIlXUtP+Mb3fMqGAmMFTLkV+Uiw0U1P162uU96tWl9ZX9+mSBTnPLaKChGbw8KGZcmWkqCioKBSwNaHorKDq6vKVmkxEHt0FMFAOJBECrJ3leUc51UBRS5xmSvoQS+2oU/1Yzt2YACDGMIwfIygip3YhVGMYRy7sQd78Rg/zvt4gvdzjQ/wQX6CJ/lJfoqfxhQ/g0N4FodxBEdxDHUc5xNo8HO2eRKnMI3TmMHzOIMXQHDYTGh2XRVwzu46pQWcrmdVTuW5wEVVQhkVjtEhRKcH9KAe0sNG0HGVd3KMjnfzHt7Lgg77golgf1DDbeh4CiE6dZiP8FE+xnWO0eGkWQWnHbgXcRYv4RxexqytnscruIBX8Rou4nXM2Ta/gXm8ibdAKoq4jkc4jC3OcGhft4PIvqbH9Jptps/r97bzDh7gQR7iYRWsboEro1wZ48o4V0jcb0dgi5SIHe6cTRJ9kKC0BW0LUUQrCXWVUKHtmGwU5CxLdsQYz1iTNCmTNhmTNS7QxdkombIiCW+RlBXT2/4o+o2SxBSHiD6XYmemyGdJOS3jpLRJieTLaaKDMnfdRfi7Mj/w2b3fMj6VYuTKUV+fS/8gzF+R6gTnZSqR65ojERfSF2760TfL1+L5hh5eJBmtbrilTtB84+zc+dnpE42JmUajS5UO+Yb4pKzchQ8sJKbaMuQclZTebm8dzPcZ6Zciv3aVljejZrP8p+9MpAU5TDdGzWZzlfEnydaYPLlleJLFcvR2R7WrfyyXrvi068eernJyvdRVzQ8g1CQFMywVS+pWRuqI1DC32fC+82fbNeIWh3maCA0g5SMpJNePe4KT7A4CCuvrA9jEw60iqnJMcXR14d1/NwA88WYksbLGUOGKjV64g23zVHHn0N2vyJsRLjK30DnxyEf7/TFAJyNcLKzv7oa1EnqdWPlEVl13tHLx+t89D5MFEtGN2x2uLIcFyK75pci7nx0e5TEej392eqm79sg7Wf43/xMbxKW6c3XaIAqmtjRHfwMAAP//AwBQSwMEFAAGAAgAAAAhAFLK5CipAQAAGwMAABAACAFkb2NQcm9wcy9hcHAueG1sIKIEASigAAEAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAnJLBbhMxEIbvSLzDyvfGmxRVKPK6QimoBxCRkvZuvLOJhde27Okq4cYVuCJVoo/Aocc+EzTvwOyumm4oJ24z8//6/XlscbqpbdZATMa7go1HOcvAaV8atyrYxfLN0UuWJVSuVNY7KNgWEjuVz5+JefQBIhpIGUW4VLA1YphynvQaapVGJDtSKh9rhdTGFfdVZTSceX1Vg0M+yfMTDhsEV0J5FPaBrE+cNvi/oaXXLV+6XG4DAUvxKgRrtEK6pXxndPTJV5i93miwgg9FQXQL0FfR4Fbmgg9bsdDKwoyCZaVsAsEfB+IcVLu0uTIxSdHgtAGNPmbJfKK1TVj2QSVocQrWqGiUQ8JqbX3T1TYkjPL+x7fd55/3X292X+4EJ0s/7sqhe1ibF3LcGag4NLYBPQoJh5BLgxbS+2quIv6DeTxk7hh64h5nkk+Os93369+317/ubp9wdpenE/86Y+broNyWhH311riP6SIs/ZlCeFjs4VAs1ipCSW+xX/x+IM5pp9G2IbO1cisoHzxPhfYbXPZ/XY5PRvlxTi88mAn++KvlHwAAAP//AwBQSwECLQAUAAYACAAAACEAQTeCz24BAAAEBQAAEwAAAAAAAAAAAAAAAAAAAAAAW0NvbnRlbnRfVHlwZXNdLnhtbFBLAQItABQABgAIAAAAIQC1VTAj9AAAAEwCAAALAAAAAAAAAAAAAAAAAKcDAABfcmVscy8ucmVsc1BLAQItABQABgAIAAAAIQCBPpSX8wAAALoCAAAaAAAAAAAAAAAAAAAAAMwGAAB4bC9fcmVscy93b3JrYm9vay54bWwucmVsc1BLAQItABQABgAIAAAAIQBAsl3p1gIAAFAFAAAPAAAAAAAAAAAAAAAAAP8IAAB4bC93b3JrYm9vay54bWxQSwECLQAUAAYACAAAACEAJ28ezewDAAAnFAAAFAAAAAAAAAAAAAAAAAACDAAAeGwvc2hhcmVkU3RyaW5ncy54bWxQSwECLQAUAAYACAAAACEAO20yS8EAAABCAQAAIwAAAAAAAAAAAAAAAAAgEAAAeGwvd29ya3NoZWV0cy9fcmVscy9zaGVldDEueG1sLnJlbHNQSwECLQAUAAYACAAAACEAQA/yOp0GAACNGgAAEwAAAAAAAAAAAAAAAAAiEQAAeGwvdGhlbWUvdGhlbWUxLnhtbFBLAQItABQABgAIAAAAIQB6VS2vsAQAAH8XAAANAAAAAAAAAAAAAAAAAPAXAAB4bC9zdHlsZXMueG1sUEsBAi0AFAAGAAgAAAAhAN9QpOQfBwAAVyIAABgAAAAAAAAAAAAAAAAAyxwAAHhsL3dvcmtzaGVldHMvc2hlZXQxLnhtbFBLAQItABQABgAIAAAAIQC6PolWRAEAAFkCAAARAAAAAAAAAAAAAAAAACAkAABkb2NQcm9wcy9jb3JlLnhtbFBLAQItABQABgAIAAAAIQAM5xIvMwQAADgSAAAnAAAAAAAAAAAAAAAAAJsmAAB4bC9wcmludGVyU2V0dGluZ3MvcHJpbnRlclNldHRpbmdzMS5iaW5QSwECLQAUAAYACAAAACEAUsrkKKkBAAAbAwAAEAAAAAAAAAAAAAAAAAATKwAAZG9jUHJvcHMvYXBwLnhtbFBLBQYAAAAADAAMACYDAADyLQAAAAA=
""".strip()


def extract_reference_receipt(duplicate_detail):
    """
    기존 중복검사 상세정보에서 기준이 된 접수번호를 추출.

    기준 차단목록 중복:
      [URL / 접수번호 7783509]

    일일파일 내부 동일 URL 중복:
      [최초번호 7783509 / URL ...]

    두 유형이 동시에 걸린 경우에는 기존 차단목록의 접수번호를 우선 사용한다.
    """
    text = "" if duplicate_detail is None else str(duplicate_detail)

    match = re.search(r"접수번호\s*([^\]\s/]+)", text)
    if match:
        return clean_id(match.group(1))

    match = re.search(r"최초번호\s*([^\]\s/]+)", text)
    if match:
        return clean_id(match.group(1))

    return ""


def clone_template_row_format(ws, source_row, target_row):
    """샘플의 데이터 행 서식을 새 행에 그대로 복사."""
    src_dim = ws.row_dimensions[source_row]
    dst_dim = ws.row_dimensions[target_row]

    dst_dim.height = src_dim.height
    dst_dim.hidden = src_dim.hidden
    dst_dim.outlineLevel = src_dim.outlineLevel

    for col in range(1, ws.max_column + 1):
        src_cell = ws.cell(source_row, col)
        dst_cell = ws.cell(target_row, col)

        if src_cell.has_style:
            dst_cell._style = copy(src_cell._style)

        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.protection = copy(src_cell.protection)

        if src_cell.number_format:
            dst_cell.number_format = src_cell.number_format


def write_upload_sample_output(excel_data, out_path):
    """
    중복 결과를 '업로드샘플(수정용)'과 동일한 엑셀 양식으로 저장.

    결과 입력 위치:
      A열 = 고유번호(현재 중복된 등록/접수번호)
      R열 = 키워드 -> '중복'
      T열 = 처리비고 -> '접수번호 XXXXXXX 중복'

    나머지 열은 샘플의 제목/순서/서식은 유지하고 값은 비워 둔다.
    """
    template_bytes = base64.b64decode(TEMPLATE_XLSX_BASE64)
    wb = load_workbook(io.BytesIO(template_bytes))
    ws = wb.worksheets[0]

    # 샘플 구조가 바뀌었는지 안전 확인
    if (
        ws["A1"].value != "고유번호"
        or ws["R1"].value != "키워드"
        or ws["T1"].value != "처리비고"
        or ws.max_column != 26
    ):
        raise ValueError(
            "내장된 업로드샘플 양식이 예상 구조(A=고유번호, R=키워드, T=처리비고, A:Z 26열)와 다릅니다."
        )

    original_max_row = ws.max_row

    # 결과가 샘플 기본 행 수보다 많으면 마지막 샘플 데이터행의 형식을 복제
    # 24행은 샘플의 빈 행이므로 23행을 데이터행 서식 기준으로 사용
    style_source_row = 23 if original_max_row >= 23 else max(2, original_max_row)
    required_last_row = max(original_max_row, len(excel_data) + 1)

    for row_num in range(original_max_row + 1, required_last_row + 1):
        clone_template_row_format(ws, style_source_row, row_num)

    # 샘플에 들어 있는 예시 데이터만 모두 지운다.
    # 제목행/글꼴/테두리/열너비/행높이/페이지설정 등은 그대로 유지.
    for row in ws.iter_rows(
        min_row=2,
        max_row=required_last_row,
        min_col=1,
        max_col=ws.max_column
    ):
        for cell in row:
            cell.value = None
            cell.hyperlink = None
            cell.comment = None

    # 결과 입력
    for row_num, item in enumerate(excel_data, start=2):
        current_id = item.get("중복된 등록/접수번호", "")
        detail = item.get("중복 상세정보", "")
        reference_id = extract_reference_receipt(detail)

        ws.cell(row_num, 1).value = current_id          # A: 고유번호
        ws.cell(row_num, 18).value = "중복"             # R: 키워드

        if reference_id:
            ws.cell(row_num, 20).value = f"접수번호 {reference_id} 중복"  # T: 처리비고
        else:
            ws.cell(row_num, 20).value = "중복"

    wb.save(out_path)
    wb.close()

def run_process(daily_path, block_path):

    # -----------------------------
    # 1. 엑셀 읽기
    # -----------------------------
    df_daily = pd.read_excel(daily_path)

    df_block = pd.read_excel(
        block_path,
        sheet_name="차단목록"
    )

    # -----------------------------
    # 2. URL 열 검색
    # -----------------------------
    def find_col(df, keywords):
        for col in df.columns:
            cleaned = str(col).replace(" ", "").lower()

            for keyword in keywords:
                if keyword.lower() in cleaned:
                    return col

        return None

    daily_url_col = find_col(
        df_daily,
        ["URL", "주소", "링크"]
    )

    block_url_col = find_col(
        df_block,
        ["URL", "주소", "링크"]
    )

    block_id_col = find_col(
        df_block,
        ["접수번호", "기존접수번호", "관리번호"]
    )

    # -----------------------------
    # 3. 일일파일 번호는 무조건 A열
    # -----------------------------
    daily_id_col = df_daily.columns[0]

    if daily_url_col is None:
        raise ValueError(
            "일일파일에서 URL 열을 찾을 수 없습니다."
        )

    if block_url_col is None:
        raise ValueError(
            "기준파일에서 URL 열을 찾을 수 없습니다."
        )

    # -----------------------------
    # 4. 비교키 생성
    # -----------------------------

    # 기준파일 ↔ 일일파일 비교용:
    # 기존대로 메인 도메인 기준
    df_daily["검사용_도메인"] = (
        df_daily[daily_url_col]
        .apply(extract_domain)
    )

    df_block["검사용_도메인"] = (
        df_block[block_url_col]
        .apply(extract_domain)
    )

    # 일일파일 내부 비교용:
    # 도메인이 아니라 하위 경로를 포함한 전체 URL 기준
    df_daily["검사용_전체URL"] = (
        df_daily[daily_url_col]
        .apply(normalize_full_url)
    )

    # -----------------------------
    # 5. 기존 차단목록 딕셔너리 생성
    # -----------------------------
    mapping_dict = {}

    for _, row in df_block.iterrows():

        dom = row["검사용_도메인"]

        if not dom:
            continue

        if (
            block_id_col is not None
            and pd.notna(row[block_id_col])
        ):
            block_id = clean_id(
                row[block_id_col]
            )
        else:
            block_id = "접수번호없음"

        if pd.notna(row[block_url_col]):
            block_url = str(
                row[block_url_col]
            ).strip()
        else:
            block_url = "URL없음"

        # 동일 도메인이 기준파일에 여러 개 있어도
        # 기존 프로그램과 동일하게 최초 데이터 하나만 저장
        if dom not in mapping_dict:
            mapping_dict[dom] = {
                "url": block_url,
                "id": block_id
            }

    # -----------------------------
    # 6. 중복 검사
    # -----------------------------

    dup_ids = []
    excel_data = []

    # ★ 중요:
    # 예전 EXE에서는 seen_daily에 '도메인'을 저장해서
    # t.me/AAA 와 t.me/BBB도 중복으로 처리했음.
    #
    # 수정본에서는 '정규화된 전체 URL'을 저장한다.
    seen_daily_urls = {}

    for _, row in df_daily.iterrows():

        dom = row["검사용_도메인"]
        full_url_key = row["검사용_전체URL"]

        if not dom:
            continue

        current_id = clean_id(
            row[daily_id_col]
        )

        current_url = (
            str(row[daily_url_col]).strip()
            if pd.notna(row[daily_url_col])
            else ""
        )

        duplicate_types = []
        duplicate_infos = []

        # =====================================
        # ① 기존 차단목록과 중복 검사
        #    → 기존대로 도메인 기준
        # =====================================

        if dom in mapping_dict:

            block_info = mapping_dict[dom]

            duplicate_types.append(
                "기존 차단목록 중복"
            )

            duplicate_infos.append(
                f"[{block_info['url']} / "
                f"접수번호 {block_info['id']}]"
            )

        # =====================================
        # ② 일일파일 내부 중복 검사
        #    → 전체 URL 기준
        # =====================================

        if full_url_key:

            if full_url_key in seen_daily_urls:

                first_info = seen_daily_urls[full_url_key]

                duplicate_types.append(
                    "일일파일 내부 동일 URL 중복"
                )

                duplicate_infos.append(
                    f"[최초번호 {first_info['id']} / "
                    f"URL {first_info['url']}]"
                )

            else:

                # 최초 등장 전체 URL 기록
                seen_daily_urls[full_url_key] = {
                    "id": current_id,
                    "url": current_url
                }

        # =====================================
        # 중복이면 결과 저장
        # =====================================

        if duplicate_types:

            if current_id:
                dup_ids.append(current_id)

            excel_data.append({

                "중복된 등록/접수번호":
                    current_id,

                "일일파일 URL":
                    current_url,

                "도메인":
                    dom,

                "중복유형":
                    " + ".join(
                        duplicate_types
                    ),

                "중복 상세정보":
                    " / ".join(
                        duplicate_infos
                    )
            })

    # -----------------------------
    # 7. 중복번호 중복제거
    # -----------------------------

    # 순서는 유지하면서 번호 중복 제거
    dup_ids = list(
        dict.fromkeys(dup_ids)
    )

    # -----------------------------
    # 8. 업로드샘플 양식으로 결과 엑셀 생성
    # -----------------------------

    dir_name = os.path.dirname(
        daily_path
    )

    base_name = os.path.splitext(
        os.path.basename(daily_path)
    )[0]

    out_path = os.path.join(
        dir_name,
        f"{base_name}_중복검사완료.xlsx"
    )

    # 첨부된 업로드샘플과 동일한 A:Z 26열 양식으로 저장
    write_upload_sample_output(
        excel_data,
        out_path
    )

    return (
        out_path,
        dup_ids,
        len(excel_data)
    )


# ==========================================
# 화면 UI
# ==========================================

class App:

    def __init__(self, root):

        self.root = root

        self.root.title(
            "저작권 URL 중복검사기 v1.5"
        )

        self.root.geometry(
            "520x500"
        )

        self.root.resizable(
            False,
            False
        )

        self.block_path = ""
        self.daily_path = ""

        # -----------------------------------
        # 기준파일
        # -----------------------------------

        tk.Label(
            root,
            text="[1] 기준 파일: 중복검사(저작권).xlsm",
            font=("맑은 고딕", 9, "bold")
        ).pack(
            anchor="w",
            padx=20,
            pady=(15, 2)
        )

        self.lbl_block = tk.Label(
            root,
            text="파일을 선택하세요...",
            fg="gray",
            anchor="w"
        )

        self.lbl_block.pack(
            fill="x",
            padx=20
        )

        tk.Button(
            root,
            text="차단목록 파일 선택",
            command=self.select_block
        ).pack(
            anchor="w",
            padx=20,
            pady=(2, 10)
        )

        # -----------------------------------
        # 일일파일
        # -----------------------------------

        tk.Label(
            root,
            text="[2] 일일 파일: 모니터/일반 내보내기 파일",
            font=("맑은 고딕", 9, "bold")
        ).pack(
            anchor="w",
            padx=20,
            pady=(5, 2)
        )

        self.lbl_daily = tk.Label(
            root,
            text="파일을 선택하세요...",
            fg="gray",
            anchor="w"
        )

        self.lbl_daily.pack(
            fill="x",
            padx=20
        )

        tk.Button(
            root,
            text="일일 파일 선택",
            command=self.select_daily
        ).pack(
            anchor="w",
            padx=20,
            pady=(2, 15)
        )

        # -----------------------------------
        # 검사버튼
        # -----------------------------------

        tk.Button(
            root,
            text="중복 검사 시작하기",
            bg="#0078D7",
            fg="white",
            font=("맑은 고딕", 11, "bold"),
            height=2,
            command=self.process
        ).pack(
            fill="x",
            padx=20
        )

        # -----------------------------------
        # 결과창
        # -----------------------------------

        tk.Label(
            root,
            text="[중복된 등록번호 / 접수번호]",
            font=("맑은 고딕", 9, "bold")
        ).pack(
            anchor="w",
            padx=20,
            pady=(15, 2)
        )

        self.result_text = (
            scrolledtext.ScrolledText(
                root,
                height=6,
                width=60,
                font=("맑은 고딕", 10)
            )
        )

        self.result_text.pack(
            padx=20,
            pady=(0, 10)
        )

        self.result_text.insert(
            tk.END,
            "검사를 완료하면 중복된 등록번호/"
            "접수번호가 표시됩니다.\n"
            "(전체 선택 후 복사해서 시스템에서 "
            "검색하세요)"
        )

    # ---------------------------------------
    # 기준파일 선택
    # ---------------------------------------

    def select_block(self):

        file = filedialog.askopenfilename(
            filetypes=[
                (
                    "Excel Files",
                    "*.xlsx *.xlsm"
                )
            ]
        )

        if file:

            self.block_path = file

            self.lbl_block.config(
                text=os.path.basename(file),
                fg="black"
            )

    # ---------------------------------------
    # 일일파일 선택
    # ---------------------------------------

    def select_daily(self):

        file = filedialog.askopenfilename(
            filetypes=[
                (
                    "Excel Files",
                    "*.xlsx *.xls"
                )
            ]
        )

        if file:

            self.daily_path = file

            self.lbl_daily.config(
                text=os.path.basename(file),
                fg="black"
            )

    # ---------------------------------------
    # 중복 검사 실행
    # ---------------------------------------

    def process(self):

        if (
            not self.block_path
            or not self.daily_path
        ):

            messagebox.showwarning(
                "경고",
                "두 파일을 모두 선택해주세요."
            )

            return

        try:

            self.result_text.delete(
                1.0,
                tk.END
            )

            self.result_text.insert(
                tk.END,
                "검사 중입니다..."
            )

            self.root.update()

            (
                out_file,
                dup_ids,
                dup_count
            ) = run_process(
                self.daily_path,
                self.block_path
            )

            self.result_text.delete(
                1.0,
                tk.END
            )

            if dup_ids:

                result_str = ",".join(
                    dup_ids
                )

                self.result_text.insert(
                    tk.END,
                    result_str
                )

                messagebox.showinfo(
                    "완료",
                    f"총 {dup_count}건의 "
                    f"중복이 발견되었습니다!\n\n"
                    f"등록/접수번호 {len(dup_ids)}개가 "
                    f"표시되었습니다.\n\n"
                    f"생성된 엑셀 파일에서 "
                    f"상세내역을 확인하세요."
                )

            else:

                self.result_text.insert(
                    tk.END,
                    "중복된 등록번호/"
                    "접수번호가 없습니다."
                )

                messagebox.showinfo(
                    "완료",
                    "검사 결과 중복건이 없습니다!"
                )

        except Exception as e:

            self.result_text.delete(
                1.0,
                tk.END
            )

            self.result_text.insert(
                tk.END,
                f"오류 발생: {str(e)}"
            )

            messagebox.showerror(
                "오류 발생",
                f"작업 중 오류가 발생했습니다:\n"
                f"{str(e)}"
            )


if __name__ == "__main__":

    root = tk.Tk()

    app = App(root)

    root.mainloop()
